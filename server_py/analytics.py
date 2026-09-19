"""
Сводная аналитика и экспорт в Excel (порт statsData/analyticsSheets из server.js).
Excel собирается через openpyxl (обычная зависимость), формат листов совпадает.
"""
from __future__ import annotations

import io

import db
from bank import BY_ID


def stats_data(stream, ac):
    if stream is None or stream == "all":
        stream = None
    else:
        try:
            stream = int(stream)
        except (TypeError, ValueError):
            stream = None

    sessions = []
    for s in db.list_sessions(stream):
        answers = s.get("answers") or 0
        correct = s.get("correct") or 0
        avg = s.get("avgScore")
        sessions.append({
            "id": s["id"], "date": s.get("date"), "stream": s.get("stream"),
            "status": s.get("status"),
            "started_at": s.get("started_at"), "finished_at": s.get("finished_at"),
            "students": s.get("students") or 0, "answers": answers, "correct": correct,
            "accuracy": round(correct / answers * 100) if answers else 0,
            "avgScore": round(avg) if avg is not None else None,
        })

    tasks = []
    for t in db.task_stats(stream):
        meta = BY_ID.get(t["taskId"])
        attempts = t.get("attempts") or 0
        correct = t.get("correct") or 0
        avg_ms = t.get("avgMs")
        tasks.append({
            "id": t["taskId"],
            "title": meta["title"] if meta else t["taskId"],
            "prompt": meta["prompt"] if meta else "",
            "topic": meta.get("topic") if meta else "",
            "type": meta["type"] if meta else "",
            "level": (meta.get("level", 1) if meta else 1),
            "attempts": attempts, "correct": correct,
            "accuracy": round(correct / attempts * 100) if attempts else 0,
            "avgMs": round(avg_ms) if avg_ms is not None else None,
        })

    students = []
    threshold = (ac or {}).get("suspectThreshold", 4) or 4
    for s in db.student_stats(stream):
        topic_map = {}
        for tr in (s.get("taskRows") or []):
            meta = BY_ID.get(tr["taskId"])
            topic = meta.get("topic") if meta else "other"
            cur = topic_map.get(topic) or {"topic": topic, "attempts": 0, "correct": 0}
            cur["attempts"] += tr["attempts"]
            cur["correct"] += tr["correct"]
            topic_map[topic] = cur
        suspicion = (s.get("tabLeaves") or 0) + 2 * (s.get("fastAnswers") or 0)
        avg_ms = s.get("avgMs")
        avg_score = s.get("avgScore")
        students.append({
            "nick": s["nick"], "icon": s["icon"], "sessions": s["sessions"],
            "answers": s["answers"], "correct": s["correct"], "accuracy": s["accuracy"],
            "avgMs": round(avg_ms) if avg_ms is not None else None,
            "avgScore": round(avg_score) if avg_score is not None else None,
            "bestPlace": s.get("bestPlace"), "tabLeaves": s.get("tabLeaves"),
            "fastAnswers": s.get("fastAnswers") or 0,
            "suspicion": suspicion, "suspect": suspicion >= threshold,
            "topics": list(topic_map.values()),
            "sessionRows": [
                {**r, "avgMs": round(r["avgMs"]) if r["avgMs"] is not None else None}
                for r in (s.get("sessionRows") or [])
            ],
        })

    attendance = db.attendance_map(stream)
    return {"sessions": sessions, "tasks": tasks, "students": students,
            "attendance": attendance, "stream": "all" if stream is None else stream,
            "backend": db.backend}


def analytics_sheets(stream, ac):
    d = stats_data(stream, ac)

    def ms(v):
        return "" if v is None else round(v / 1000, 1)

    topic_name = {"asd": "АиСД", "graphs": "Графы"}
    type_name = {"choice": "выбор", "graph": "построение"}

    students = [["Логин", "Сессий", "Ответов", "Верных", "Точность %", "Ср. балл",
                 "Лучшее место", "Уходы со вкладки", "Быстрые ответы", "Индекс подозр.", "Подозрит."]]
    for s in d["students"]:
        students.append([
            s["nick"], s["sessions"], s["answers"], s["correct"], s["accuracy"],
            s["avgScore"] if s["avgScore"] is not None else "",
            s["bestPlace"] if s["bestPlace"] is not None else "",
            s["tabLeaves"] or 0, s["fastAnswers"] or 0, s["suspicion"] or 0,
            "да" if s["suspect"] else "",
        ])

    tasks = [["ID", "Задача", "Тема", "Тип", "Сложность", "Попыток", "Верных", "Точность %", "Ср. время, с"]]
    for t in d["tasks"]:
        tasks.append([
            t["id"], t["title"], topic_name.get(t["topic"], t["topic"]),
            type_name.get(t["type"], t["type"]), t["level"],
            t["attempts"], t["correct"], t["accuracy"], ms(t["avgMs"]),
        ])

    sessions = [["#", "Дата", "Поток", "Статус", "Участников", "Ответов", "Верных", "Точность %", "Ср. балл"]]
    for s in d["sessions"]:
        sessions.append([
            s["id"], s["date"] or "", s["stream"] if s["stream"] is not None else "", s["status"],
            s["students"], s["answers"], s["correct"], s["accuracy"],
            s["avgScore"] if s["avgScore"] is not None else "",
        ])

    cols = sorted(d["sessions"], key=lambda c: c["id"])
    header = ["Логин"] + [
        f"#{c['id']} {c['date'] or ''}" + (f" (п{c['stream']})" if c["stream"] is not None else "")
        for c in cols
    ] + ["Всего"]
    att = [header]
    for a in d["attendance"]:
        sset = set(a.get("sessionIds") or [])
        att.append([a["nick"]] + ["✓" if c["id"] in sset else "" for c in cols] + [len(sset)])

    return [
        {"name": "Логины", "rows": students},
        {"name": "Задачи", "rows": tasks},
        {"name": "Сессии", "rows": sessions},
        {"name": "Посещаемость", "rows": att},
    ]


def build_xlsx(sheets):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    wb.remove(wb.active)
    bold = Font(bold=True)
    for sh in sheets:
        ws = wb.create_sheet(title=sh["name"][:31])
        for ri, row in enumerate(sh["rows"]):
            for ci, val in enumerate(row):
                cell = ws.cell(row=ri + 1, column=ci + 1, value=val)
                if ri == 0:
                    cell.font = bold
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
